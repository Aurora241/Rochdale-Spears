import openpyxl,datetime as dt,collections,sys
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
F='Times New Roman';t=Side(style='thin');bd=Border(left=t,right=t,top=t,bottom=t)
red=PatternFill('solid',fgColor='C00000');hf=Font(name=F,bold=True,color='FFFFFF');ctr=Alignment(horizontal='center',vertical='center',wrap_text=True)
band=PatternFill('solid',fgColor='1F4E78');grp=PatternFill('solid',fgColor='DDEBF7');gry=PatternFill('solid',fgColor='D9D9D9')
ORDER=['P20','P26','P22','A2']        # thứ tự mã cố định trong mọi bảng
AREAS=['Assembly','Machining','Raw Cutting']
WD=['Thứ Hai','Thứ Ba','Thứ Tư','Thứ Năm','Thứ Sáu','Thứ Bảy','Chủ nhật']
def tmin(x): return x.hour*60+x.minute if hasattr(x,'hour') else 0
def sheet(wb,title,rows):
    """rows: (date,code,notes,process,dept,from,to,hc)"""
    ws=wb.create_sheet(title)
    ws['A1']=f'BẢNG NHẬP LIỆU - {title.upper()}';ws['A1'].font=Font(name=F,bold=True,size=14)
    ws['A2']='Thứ tự cố định: theo ngày, trong mỗi ngày theo mã P20, P26, P22, A2. Nhập đúng thứ tự cột như form: LossTime Code | Notes | Process | Dept In Charge | From | To | Worker Qty. Nhập xong dòng nào thì đánh dấu X cột "Đã nhập".'
    ws['A2'].font=Font(name=F,italic=True)
    H=['Ngày (Loss Date)','LossTime Code','Notes','Process','Dept In Charge','From','To','Worker Qty','Giờ công (tham khảo)','Đã nhập']
    for i,h in enumerate(H,1):
        c=ws.cell(4,i,h);c.fill=red if i<=8 else gry;c.font=hf if i<=8 else Font(name=F,bold=True);c.alignment=ctr;c.border=bd
    ws.row_dimensions[4].height=32
    r=5;days=sorted({x[0] for x in rows});first=r
    for d in days:
        dr=[x for x in rows if x[0]==d]
        dr.sort(key=lambda x:(ORDER.index(x[1]) if x[1] in ORDER else 9,tmin(x[5]),str(x[2])))
        ws.cell(r,1,f'{WD[d.weekday()]} {d:%d/%m/%Y}: {len(dr)} dòng');
        for c in range(1,11):
            x=ws.cell(r,c);x.fill=band;x.font=Font(name=F,bold=True,color='FFFFFF')
        bandrow=r;r+=1;ds=r
        prev=None
        for x in dr:
            if prev is not None and x[1]!=prev: pass
            prev=x[1]
            row=[x[0],x[1],x[2],x[3],x[4],x[5],x[6],x[7],f'=(G{r}-F{r})*24*H{r}',None]
            for c,v in enumerate(row,1):
                cc=ws.cell(r,c,v);cc.font=Font(name=F);cc.border=bd;cc.alignment=Alignment(horizontal='left' if c==3 else 'center',vertical='center',wrap_text=c==3)
                if x[1] in('P22','A2'): cc.fill=grp
            ws.cell(r,1).number_format='dd/mm/yyyy';ws.cell(r,6).number_format=ws.cell(r,7).number_format='hh:mm';ws.cell(r,9).number_format='0.0'
            r+=1
        ws.cell(bandrow,9,f'=SUM(I{ds}:I{r-1})').number_format='0.0'
        ws.cell(bandrow,9).font=Font(name=F,bold=True,color='FFFFFF')
        r+=1
    ws.cell(r,8,'Tổng').font=Font(name=F,bold=True)
    ws.cell(r,9,f'=SUMIF(A{first}:A{r-1},">0",I{first}:I{r-1})').font=Font(name=F,bold=True);ws.cell(r,9).number_format='#,##0.0'
    for col,w in zip('ABCDEFGHIJ',[22,12,46,14,18,8,8,10,12,9]): ws.column_dimensions[col].width=w
    ws.freeze_panes='A5'
    return ws
def build(rows,out,fixes=None,extra=None):
    wb=openpyxl.Workbook();wb.remove(wb.active)
    for a in AREAS:
        sheet(wb,a,[x for x in rows if x[3]==a])
    wb.save(out)
if __name__=='__main__':
    # dùng: python byarea.py Nhap_lieu_xxx.xlsx ngay_dau ngay_cuoi out.xlsx  (ngày dạng dd-mm)
    src,d0,d1,out=sys.argv[1:5]
    wb=openpyxl.load_workbook(src,data_only=True)
    rows=[]
    for ws in wb.worksheets:
        if not ws.title.startswith('Nhập '): continue
        dd,mm=ws.title[5:].split('-');d=dt.datetime(2026,int(mm),int(dd))
        if not(dt.datetime(2026,int(d0[3:]),int(d0[:2]))<=d<=dt.datetime(2026,int(d1[3:]),int(d1[:2]))): continue
        for r in ws.iter_rows(min_row=2,values_only=True):
            if r[0]: rows.append((d,r[0],r[1],r[2],r[3],r[4],r[5],r[6]))
    build(rows,out)
