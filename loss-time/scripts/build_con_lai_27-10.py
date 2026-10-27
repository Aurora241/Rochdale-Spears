import openpyxl,datetime as dt
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
import byarea as B
T=lambda h,m=0:dt.time(h,m);D=lambda d:dt.datetime(2026,10,d)
F=B.F;bd=B.bd
adds=[
(D(24),'P26','Ghi nhận hàng giao kho','Machining','MACHINING',T(7,30),T(15,30),1),
(D(24),'A2','HC mới- 50% hiệu suất- 2 tuần-23/10','Raw Cutting','HR/ HSE',T(7,30),T(12,0),1),
]
wb=openpyxl.Workbook();wb.remove(wb.active)
for a in B.AREAS: B.sheet(wb,a,[x for x in adds if x[3]==a])
ws=wb.create_sheet('Sửa theo bộ phận')
ws['A1']='DÒNG ĐÃ NHẬP CẦN SỬA (đối chiếu file xuất 27/10, 19/10-25/10). Xếp theo bộ phận, ngày, mã.';ws['A1'].font=Font(name=F,bold=True,size=13)
H=['Bộ phận','Ngày','Mã','Nhận ra dòng bằng','Ô cần sửa','Đang có','Sửa thành','Giờ công thay đổi','Đã làm']
for i,h in enumerate(H,1):
    c=ws.cell(3,i,h);c.fill=B.red;c.font=B.hf;c.alignment=B.ctr;c.border=bd
fx=[
('Assembly',D(21),'A2','Notes "HC mới- 50% hiệu suất- 2 tuần-09/10", Assembly, 7:30-12:00, 6 người: đang có 2 dòng giống nhau. Sửa 1 trong 2 dòng thành Raw Cutting 2 người','Process và Worker Qty','Assembly, 6','Raw Cutting, 2',-18),
('Assembly',D(25),'A2','Notes "...2 tuần-13/10", Assembly, 7:30-12:00','Worker Qty','4','2',-9),
('Assembly',D(25),'A2','Notes "...2 tuần-15/10", Assembly, 7:30-12:00','Worker Qty','2','4',9),
('Assembly',D(25),'A2','Notes "...2 tuần-19/10", Assembly, 7:30-12:00','Worker Qty','1','2',4.5),
('Machining',D(24),'P20','Vệ sinh khu vực máy móc MAC FINAL (Machining, 7:00-7:30, 7 người)','Dept In Charge','RAW CUTTING','MACHINING',0),
('Machining',D(24),'P20','Vệ sinh khu vực máy móc PRE MAC (Machining, 7:00-7:30, 12 người)','Dept In Charge','RAW CUTTING','MACHINING',0),
('Machining',D(24),'P26','Shop Floor data MAC FINAL (Machining, 7:30-15:30, 1 người)','Dept In Charge','RAW CUTTING','MACHINING',0),
('Machining',D(24),'P26','Shop Floor data- STOCK ROOM (Machining, 7:30-15:30, 1 người)','Dept In Charge','RAW CUTTING','MACHINING',0),
]
fx.sort(key=lambda x:(B.AREAS.index(x[0]),x[1],B.ORDER.index(x[2])))
for k,it in enumerate(fx):
    r=4+k
    for c,v in enumerate(list(it)+[None],1):
        x=ws.cell(r,c,v);x.font=Font(name=F);x.border=bd;x.alignment=Alignment(horizontal='left' if c==4 else 'center',vertical='center',wrap_text=c==4)
    ws.cell(r,2).number_format='dd/mm/yyyy'
for col,w in zip('ABCDEFGHI',[14,12,7,70,22,16,16,14,9]): ws.column_dimensions[col].width=w
s4=wb.create_sheet('S32 Raw Cutting còn lại')
s4['A1']='S32 RAW CUTTING CÒN LẠI (hệ thống 27/10 vẫn chưa sửa). Assembly bạn tự nhập.';s4['A1'].font=Font(name=F,bold=True,size=13)
hh=['Việc cần làm','Ngày','Notes','From','To','Worker Qty','Dept In Charge','Giờ công','Ghi chú','Đã làm']
for i,h in enumerate(hh,1):
    c=s4.cell(3,i,h);c.fill=B.red;c.font=B.hf;c.alignment=B.ctr;c.border=bd
uv='Không có bomtime chuyền UV'
s=[('NHẬP THÊM',D(19),uv,T(7,30),T(15,30),12,'FINISHING',96,'UV ca thường'),
('NHẬP THÊM',D(19),uv,T(15,30),T(19,30),10,'FINISHING',40,'UV tăng ca'),
('NHẬP THÊM',D(20),uv,T(7,30),T(15,30),12,'FINISHING',96,'UV ca thường'),
('NHẬP THÊM',D(20),uv,T(15,30),T(19,30),10,'FINISHING',40,'UV tăng ca'),
('NHẬP THÊM',D(21),uv,T(7,30),T(15,30),12,'FINISHING',96,'UV ca thường (tăng ca đã có)'),
('XÓA',D(24),uv+' (dòng 12 người)',T(7,30),T(15,30),12,'FINISHING',-96,'Giữ dòng 11 người (88 giờ), khớp báo cáo UV'),
('SỬA NGÀY',D(18),'Beam saw A05-S01-1 (AR1113F1330) 18:00-18:02',T(18,0),T(18,2),3,'MACHINING',0,'Loss Date 18/10 đổi thành 19/10'),
('SỬA NGÀY',D(18),'Beam saw A05-S01-2 (AR1113F1330) 18:02-18:04',T(18,2),T(18,4),3,'MACHINING',0,'Loss Date 18/10 đổi thành 19/10'),
('SỬA NGÀY',D(18),'Beam saw A1-S02 (SH20-1029-1471-41191) 18:04-18:07',T(18,4),T(18,7),3,'MACHINING',0,'Loss Date 18/10 đổi thành 19/10'),
('ĐIỀN BỘ PHẬN',D(20),'Beam saw CZ567F08 18:10-18:15',T(18,10),T(18,15),3,'(đang trống)',0,'Đề xuất OTHERS (INVESTIGATE)'),
('ĐIỀN BỘ PHẬN',D(20),'Beam saw CZ567F09 18:15-18:18',T(18,15),T(18,18),3,'(đang trống)',0,'Đề xuất OTHERS (INVESTIGATE)'),
('ĐIỀN BỘ PHẬN',D(20),'Veneer walnut cắt ghép (bs lại) RH21-2184-492 10:00-11:15',T(10,0),T(11,15),2,'(đang trống)',0,'Nguồn ghi "bs lại", cần hỏi bộ phận đúng')]
for k,row in enumerate(s):
    r=4+k
    for c,v in enumerate(list(row)+[None],1):
        x=s4.cell(r,c,v);x.font=Font(name=F);x.border=bd;x.alignment=Alignment(horizontal='left' if c in(3,9) else 'center',vertical='center',wrap_text=c in(3,9))
    s4.cell(r,2).number_format='dd/mm/yyyy';s4.cell(r,4).number_format=s4.cell(r,5).number_format='hh:mm'
for col,w in zip('ABCDEFGHIJ',[16,12,56,8,8,10,16,10,44,9]): s4.column_dimensions[col].width=w
wb.save('Nhap_lieu_con_lai_theo_bo_phan_27-10.xlsx')
