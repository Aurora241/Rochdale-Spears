import openpyxl,datetime as dt
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
src=openpyxl.load_workbook('Loss_time_05-10_den_30-10.xlsx',data_only=True)
F='Times New Roman';t=Side(style='thin');bd=Border(left=t,right=t,top=t,bottom=t)
H=['LossTime Code','Notes','Process','Dept In Charge','From','To','Worker Qty']
wb=openpyxl.Workbook();wb.remove(wb.active)
def rows(ws): return [[c.value for c in r] for r in ws.iter_rows(min_row=3,max_col=17) if r[6].value]
PR={a:'Raw Cutting' for a in ('RC - WOOD','RC - PANEL','RC','Raw Cutting Panel','RC WOOD')}
PR['Machining']='Machining'
DP={'RC - PANEL':'RAW CUTTING','RC WOOD':'RAW CUTTING','RC - WOOD':'RAW CUTTING','Machining':'MACHINING','Assembly':'ASSEMBLY','HR/HSE':'HR/ HSE'}
exp={}
for nm in ('05-10','06-10','07-10','08-10','09-10','10-10','12-10','13-10','14-10','15-10','16-10','17-10','19-10','20-10','21-10','22-10','23-10','24-10','25-10','26-10','27-10','28-10','29-10','30-10'):
    s=src[f'Loss time {nm}'];ws=wb.create_sheet(f'Nhập {nm}');data=rows(s)
    for i,h in enumerate(H+['(Tham khảo - không nhập) Ghi chú'],1):
        c=ws.cell(1,i,h);c.font=Font(name=F,bold=True,color='FFFFFF' if i<=7 else '000000');c.border=bd
        c.fill=PatternFill('solid',fgColor='C00000' if i<=7 else 'D9D9D9');c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    out=[]
    for k,r in enumerate(data,2):
        # r idx: 6 area,7 code,9 dept,10 lido,11 ghichu,12 from,13 to,15 hc
        vals=[r[7],r[10],PR.get(r[6],r[6]),DP.get(r[9],r[9]),r[12],r[13],r[15],r[11]]
        out.append(vals[:7])
        for i,v in enumerate(vals,1):
            c=ws.cell(k,i,v);c.font=Font(name=F,color='000000' if i<=7 else '7F7F7F');c.border=bd
            c.alignment=Alignment(horizontal='left' if i in(2,8) else 'center',vertical='center',wrap_text=i in(2,8))
            if i in(5,6):c.number_format='hh:mm'
    exp[nm]=out
    for col,w in zip('ABCDEFGH',[14,44,20,18,9,9,12,32]): ws.column_dimensions[col].width=w
    ws.row_dimensions[1].height=32;ws.freeze_panes='A2'
g=wb.create_sheet('Hướng dẫn')
lines=['Cột nhập theo thứ tự từ trái sang phải đúng như trên form: LossTime Code | Notes | Process | Dept In Charge | From | To | Worker Qty.',
'Các cột Loss Date, LossTime Name, TotalLoss Hour do form tự điền, không nhập.',
'Mỗi sheet "Nhập dd-mm" dùng cho đúng ngày đó (05/10 đến 30/10 năm 2026). Mỗi sheet nhập từ trên xuống dưới, có thể copy cả vùng A2:G.. rồi dán vào lưới.',
'Notes = cột "Lí Do" trong sheet Total loss time. Process = cột "Khu vực phát sinh/bộ phận hỗ trợ". Dept In Charge = cột "Chart / bộ phận chịu trách nhiệm". Đây là cách ghép tôi suy ra, cần bạn xác nhận.',
'Cột H (xám) là Ghi chú cũ (Recorder/Supporter hoặc số ngày A2), chỉ để tham khảo vì form không có ô tương ứng.',
'Process và Dept In Charge đã đổi sang đúng tên trong dropdown của hệ thống (Raw Cutting, MACHINING, ASSEMBLY, RAW CUTTING, HR/ HSE).']
for i,tx in enumerate(lines,1):
    g.cell(i,1,tx).font=Font(name=F)
g.column_dimensions['A'].width=140
wb.save('Nhap_lieu_05-10_den_30-10.xlsx')
# verify
chk=openpyxl.load_workbook('Nhap_lieu_05-10_den_30-10.xlsx')
for nm in exp:
    ws=chk[f'Nhập {nm}']
    got=[[c.value for c in r][:7] for r in ws.iter_rows(min_row=2,max_col=7) if r[0].value]
    print(nm,len(got),got==exp[nm])
