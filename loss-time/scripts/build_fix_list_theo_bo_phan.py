import openpyxl,datetime as dt
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
import byarea as B
src=openpyxl.load_workbook('Danh_sach_sua_them_19-10_den_26-10.xlsx',data_only=True)
rows=[]
for r in src['1. Nhập thêm'].iter_rows(min_row=4,values_only=True):
    if r[0] and isinstance(r[0],dt.datetime) and r[1]: rows.append((r[0],r[1],r[2],r[3],r[4],r[5],r[6],r[7]))
wb=openpyxl.Workbook();wb.remove(wb.active)
for a in B.AREAS: B.sheet(wb,a,[x for x in rows if x[3]==a])
F=B.F;bd=B.bd
# sửa / xóa theo bộ phận
ws=wb.create_sheet('Sửa-Xóa theo bộ phận')
ws['A1']='DÒNG ĐÃ NHẬP CẦN SỬA HOẶC XÓA (xếp theo bộ phận, ngày, mã)';ws['A1'].font=Font(name=F,bold=True,size=13)
H=['Bộ phận','Ngày','Việc','Mã','Nhận ra dòng bằng','Ô cần sửa','Đang có','Sửa thành','Giờ công thay đổi','Đã làm']
for i,h in enumerate(H,1):
    c=ws.cell(3,i,h);c.fill=B.red;c.font=B.hf;c.alignment=B.ctr;c.border=bd
area_of=lambda txt:'Assembly' if 'Assembly' in txt and 'Raw Cutting, 7:30' not in txt else ('Raw Cutting' if 'Raw Cutting, 7:30' in txt else 'Machining')
items=[]
for r in src['2. Sửa dòng đã nhập'].iter_rows(min_row=4,values_only=True):
    if r[0] and isinstance(r[0],dt.datetime):
        txt=str(r[2]);a='Raw Cutting' if 'Nhận hàng gỗ' in txt else ('Machining' if 'MAC' in txt else 'Assembly')
        items.append((a,r[0],'SỬA',r[1],r[2],r[3],r[4],r[5],r[7]))
for r in src['3. Xóa dòng trùng'].iter_rows(min_row=4,values_only=True):
    if r[0] and isinstance(r[0],dt.datetime): items.append(('Assembly',r[0],'XÓA',r[1],r[2],'xóa '+str(r[3])+' dòng','','',-r[4]))
items.sort(key=lambda x:(B.AREAS.index(x[0]),x[1],B.ORDER.index(x[3]) if x[3] in B.ORDER else 9))
for k,it in enumerate(items):
    r=4+k
    for c,v in enumerate(list(it)+[None],1):
        x=ws.cell(r,c,v);x.font=Font(name=F);x.border=bd;x.alignment=Alignment(horizontal='left' if c==5 else 'center',vertical='center',wrap_text=c==5)
    ws.cell(r,2).number_format='dd/mm/yyyy'
for col,w in zip('ABCDEFGHIJ',[14,12,8,7,56,16,16,16,12,9]): ws.column_dimensions[col].width=w
# S32 khu vực khác
s4=wb.create_sheet('S32 Raw Cutting')
s4['A1']='S32 CỦA RAW CUTTING (UV, Beam saw, Veneer). Assembly (thợ sơn, cào cước, bắn cát) bạn nhập khi cần.';s4['A1'].font=Font(name=F,bold=True,size=13)
hh=['Việc cần làm','Ngày','LossTime Code','Notes','Process','Dept In Charge','From','To','Worker Qty','Giờ công','Ghi chú','Đã làm']
for i,h in enumerate(hh,1):
    c=s4.cell(3,i,h);c.fill=B.red;c.font=B.hf;c.alignment=B.ctr;c.border=bd
k=0
for r in src['4. S32 khu vực khác'].iter_rows(min_row=4,values_only=True):
    if not r[0]: continue
    rr=4+k
    for c,v in enumerate(list(r[:11])+[None],1):
        x=s4.cell(rr,c,v);x.font=Font(name=F);x.border=bd;x.alignment=Alignment(horizontal='left' if c in(4,11) else 'center',vertical='center',wrap_text=c in(4,11))
    s4.cell(rr,2).number_format='dd/mm/yyyy';s4.cell(rr,7).number_format=s4.cell(rr,8).number_format='hh:mm'
    k+=1
for col,w in zip('ABCDEFGHIJKL',[16,12,10,50,12,18,8,8,10,10,60,9]): s4.column_dimensions[col].width=w
wb.save('Nhap_lieu_theo_bo_phan_19-10_den_26-10.xlsx');print('adds',len(rows),'items',len(items),'s32',k)
