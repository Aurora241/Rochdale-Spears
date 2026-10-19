import openpyxl,datetime as dt
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter as L
F='Times New Roman';t=Side(style='thin');bd=Border(left=t,right=t,top=t,bottom=t)
red=PatternFill('solid',fgColor='C00000');hf=Font(name=F,bold=True,color='FFFFFF');ctr=Alignment(horizontal='center',vertical='center',wrap_text=True)
yel=PatternFill('solid',fgColor='FFF2CC')
T=lambda h,m:dt.time(h,m)
D=lambda d:dt.datetime(2026,10,d)
SON='Không có bomtime sơn lòng trong'
rows=[ # (day, item, qty, area, dept, lido, ghichu, from, to, hc, flag)
 (12,None,None,'Assembly','Finishing',SON,'Thợ sơn lòng trong tủ - hỗ trợ Fin',T(7,30),T(19,30),12,''),
 (13,None,None,'Assembly','Finishing',SON,'Thợ sơn lòng trong tủ - hỗ trợ Fin',T(7,30),T(19,30),12,''),
 (14,None,None,'Assembly','Finishing',SON,'Thợ sơn lòng trong tủ - hỗ trợ Fin',T(7,30),T(19,30),12,''),
 (15,None,None,'Assembly','Finishing',SON,'Thợ sơn lòng trong tủ - hỗ trợ Fin',T(7,30),T(19,30),12,''),
 (15,'RH2252',734,'Assembly','Subcon','Bắn cát hỗ trợ Subcon Đại Xuyên (chân)','Line bắn cát - hỗ trợ Supcon Đại Xuyên',T(7,0),T(15,0),4,'A'),
 (16,None,None,'Assembly','Finishing',SON,'Thợ sơn lòng trong tủ - hỗ trợ Fin',T(7,30),T(19,30),12,''),
 (16,'RH893',None,'Assembly','Machining','Cào cước hỗ trợ MAC','Line 8 - hỗ trợ MAC',T(15,0),T(20,0),1,'P'),
 (16,'RH5061',None,'Assembly','Machining','Cào cước hỗ trợ MAC','Line 7 - hỗ trợ MAC',T(8,0),T(12,0),1,''),
 (17,None,None,'Assembly','Finishing',SON,'Thợ sơn lòng trong tủ - hỗ trợ Fin',T(7,30),T(15,30),12,'P'),
]
wb=openpyxl.Workbook()
ws=wb.active;ws.title='S32 tuần 42'
ws['A1']='LOSS TIME MÃ S32 - KHU VỰC ASSEMBLY - TUẦN 42 (12/10 đến 18/10/2026)';ws['A1'].font=Font(name=F,bold=True,size=16)
H=['Week','Date \n(Ngày)',"Item's\nMã hàng",'Code Máy Hư',"Q'ty",'Số NCR/ME','Khu vực phát sinh/bộ phận hỗ trợ','Code','Name code','Chart / bộ phận chịu trách nhiệm','Lí Do','Ghi chú','From','To','LossTime (Hr)','HC','Loss time labor (Hrs)']
for i,h in enumerate(H,1):
    c=ws.cell(2,i,h);c.fill=red;c.font=hf;c.alignment=ctr;c.border=bd
ws.row_dimensions[2].height=37
for k,(d,item,qty,area,dept,lido,gc,fr,to,hc,flag) in enumerate(rows):
    r=3+k
    vals=[f'="W"&WEEKNUM(B{r},21)',D(d),item,None,qty,None,area,'S32','Other Support of Production',dept,lido,gc,fr,to,f'=(N{r}-M{r})*24',hc,f'=P{r}*O{r}']
    for c,v in enumerate(vals,1):
        x=ws.cell(r,c,v);x.font=Font(name=F);x.border=bd;x.alignment=Alignment(horizontal='center',vertical='center',wrap_text=c in(9,11,12))
        if flag: x.fill=yel
    ws.cell(r,2).number_format='d-mmm';ws.cell(r,13).number_format=ws.cell(r,14).number_format='h:mm;@'
    ws.cell(r,15).number_format=ws.cell(r,17).number_format='0.0'
e=2+len(rows)
ws.cell(e+1,16,'Tổng').font=Font(name=F,bold=True);ws.cell(e+1,17,f'=SUM(Q3:Q{e})').font=Font(name=F,bold=True);ws.cell(e+1,17).number_format='0.0'
ws.cell(e+3,1,'Màu vàng = giờ cần xác nhận: (A) bắn cát ghi 7:00-15:00 để đúng 8 giờ/người như báo cáo (7:00-16:00 sẽ thành 36 giờ); (P) giờ dự kiến chưa chắc đã xảy ra (cào cước 15:00-20:00 ngày 16/10, thợ sơn thứ Bảy 17/10).').font=Font(name=F,italic=True)
for col,w in zip('ABCDEFGHIJKLMNOPQ',[7,9,10,10,8,10,16,8,26,16,40,36,9,9,10,8,15]): ws.column_dimensions[col].width=w
ws.freeze_panes='A3';ws.auto_filter.ref=f'A2:Q{e}'

# check sheet
ck=wb.create_sheet('Đối chiếu báo cáo')
ck['A1']='ĐỐI CHIẾU GIỜ CÔNG VỚI BÁO CÁO GỐC';ck['A1'].font=Font(name=F,bold=True,size=14)
hh=['Việc','Giờ công theo báo cáo','Giờ công trong bảng S32','Chênh lệch','Kết quả']
for i,h in enumerate(hh,1):
    c=ck.cell(3,i,h);c.fill=red;c.font=hf;c.alignment=ctr;c.border=bd
sh="'S32 tuần 42'"
items=[('Thợ sơn lòng trong tủ (12 người)',816,'Không có bomtime sơn lòng trong'),('Cào cước (2 dòng, 16/10)',9,'Cào cước hỗ trợ MAC'),('Bắn cát Supcon Đại Xuyên (4 người, 15/10)',32,'Bắn cát hỗ trợ Subcon Đại Xuyên (chân)')]
for k,(nm,rep,lido) in enumerate(items):
    r=4+k
    ck.cell(r,1,nm);ck.cell(r,2,rep)
    ck.cell(r,3,f"=SUMIFS({sh}!$Q$3:$Q${e},{sh}!$K$3:$K${e},\"{lido}\")")
    ck.cell(r,4,f'=C{r}-B{r}');ck.cell(r,5,f'=IF(ABS(D{r})<0.001,"Khớp","Lệch")')
r=7;ck.cell(r,1,'Tổng')
for c,col in ((2,'B'),(3,'C'),(4,'D')): ck.cell(r,c,f'=SUM({col}4:{col}6)')
ck.cell(r,5,'=IF(ABS(D7)<0.001,"Khớp","Lệch")')
for rr in range(4,8):
    for c in range(1,6):
        x=ck.cell(rr,c);x.font=Font(name=F,bold=(rr==7));x.border=bd;x.alignment=Alignment(horizontal='left' if c==1 else 'center',vertical='center',wrap_text=True)
        if c in(2,3,4): x.number_format='#,##0.0'
for col,w in zip('ABCDE',[46,22,24,14,12]): ck.column_dimensions[col].width=w
ck.cell(9,1,'Số liệu "theo báo cáo" lấy từ 3 báo cáo của khu vực Assy gửi cho tuần 42 (816 + 9 + 32 = 857 giờ công).').font=Font(name=F,italic=True)

# input sheet
ip=wb.create_sheet('Nhập liệu S32')
hi=['Ngày (chọn trên form)','LossTime Code','Notes','Process','Dept In Charge','From','To','Worker Qty','(Tham khảo) Mã hàng','(Tham khảo) Số lượng','(Tham khảo) Ghi chú']
for i,h in enumerate(hi,1):
    c=ip.cell(1,i,h);c.fill=red if 2<=i<=8 else PatternFill('solid',fgColor='7F7F7F');c.font=hf;c.alignment=ctr;c.border=bd
ip.row_dimensions[1].height=34
for k,(d,item,qty,area,dept,lido,gc,fr,to,hc,flag) in enumerate(rows):
    r=2+k
    vals=[D(d),'S32',lido,area,dept,fr,to,hc,item,qty,gc]
    for c,v in enumerate(vals,1):
        x=ip.cell(r,c,v);x.font=Font(name=F,color='000000' if c<=8 else '7F7F7F');x.border=bd;x.alignment=Alignment(horizontal='left' if c in(3,11) else 'center',vertical='center',wrap_text=c in(3,11))
        if flag and c<=8: x.fill=yel
    ip.cell(r,1).number_format='dd/mm/yyyy';ip.cell(r,6).number_format=ip.cell(r,7).number_format='hh:mm'
for col,w in zip('ABCDEFGHIJK',[14,10,42,12,14,8,8,10,14,12,38]): ip.column_dimensions[col].width=w
ip.freeze_panes='A2'
ip.cell(len(rows)+3,1,'Nhập từng dòng theo ngày ở cột A (Loss Date trên form). Cột 2 đến 8 là các ô cần nhập theo thứ tự từ trái sang phải. Màu vàng = cần xác nhận giờ.').font=Font(name=F,italic=True)

# explanation sheet
gx=wb.create_sheet('Giải thích')
lines=[('MÃ S32 LÀ GÌ',True),
('S32 = Other Support of Production (hỗ trợ khác từ sản xuất): người của một khu vực làm việc cho bộ phận khác hoặc nhà thầu, nên giờ đó tính là loss time của khu vực cho mượn người. Theo bảng mã, nhóm là Support, bộ phận dẫn dắt cải tiến là Production.',False),
('',False),
('BA BÁO CÁO TRONG YÊU CẦU',True),
('1. Thợ sơn lòng trong tủ (Assembly hỗ trợ Fin/Finishing): 12 người, thứ Hai đến thứ Sáu 12 giờ, thứ Bảy 8 giờ, CN nghỉ. Mỗi ngày một dòng HC 12, như cách đã ghi tuần W23 (tháng 6) trong file.',False),
('2. Cào cước (Assembly hỗ trợ MAC = Machining): ngày 16/10, 2 dòng, mỗi dòng 1 người, có mã hàng RH893 và RH5061.',False),
('3. Bắn cát hàng Supcon (Assembly hỗ trợ Supcon Đại Xuyên): ngày 15/10, 4 người, mã hàng RH2252, số lượng 734 (chân).',False),
('',False),
('CÁCH ĐỌC SHEET "S32 tuần 42"',True),
('Mỗi dòng theo đúng cấu trúc sheet Total loss time: Week (tự tính), Date, Mã hàng, Q\'ty, Khu vực (Assembly), Code S32, bộ phận chịu trách nhiệm (Finishing, Machining hoặc Subcon), Lí do, From, To, HC. Giờ công = (To - From) × 24 × HC.',False),
('',False),
('ĐIỂM CẦN XÁC NHẬN VỚI NGƯỜI GỬI (các ô màu vàng)',True),
('- Giờ bắn cát: báo cáo ghi 7:00-16:00 nhưng 8 giờ/người (tổng 32). Bảng ghi 7:00-15:00 để ra đúng 32; nếu ghi 7:00-16:00 thì thành 36 giờ.',False),
('- Giờ dự kiến: cào cước 15:00-20:00 ngày 16/10 và thợ sơn thứ Bảy 17/10 có thể chưa xảy ra lúc báo cáo được gửi; cần xác nhận số thực tế.',False),
('- Giờ thợ sơn: ghi 7:30-19:30 (12 giờ) theo tiền lệ W23; nếu thực tế 8:00-20:00 thì chỉ cần đổi From/To, giờ công không đổi.',False),
('- Mã S32 và bộ phận chịu trách nhiệm (Finishing, Machining, Subcon) suy ra từ tiền lệ trong file, nên hỏi production engineer để chắc chắn.',False),
('- Tên: báo cáo ghi "Supcon", file dùng "Subcon"; bảng dùng "Subcon" cho bộ phận.',False)]
gx.column_dimensions['A'].width=130
for i,(tx,b) in enumerate(lines,1):
    c=gx.cell(i,1,tx);c.font=Font(name=F,bold=b,size=12 if b else 11);c.alignment=Alignment(wrap_text=True,vertical='top')
wb.save('Loss_time_S32_tuan42.xlsx');print('ok',len(rows))
