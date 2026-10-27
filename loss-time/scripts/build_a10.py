import openpyxl,datetime as dt
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
import byarea as B
from openpyxl.styles import PatternFill
T=lambda h,m:dt.time(h,m)
D=lambda d:dt.datetime(2026,10,d)
# (ngày, khu vực báo cáo, mã máy, tên máy, hiện tượng, từ, đến, giờ báo cáo, cột số cuối báo cáo)
rep=[
(19,'Machining','CNCR02','Máy CNC','Lỗi nút bấm',T(7,43),T(8,30),0.78,24),
(20,'RC','','Máy bơm sơn chuyền UV','Máy bơm sơn không hoạt động',T(9,1),T(9,58),0.95,12),
(19,'Machining','EDB01','Máy dán cạnh thẳng tự động','Lỗi hộp số',T(18,45),T(19,45),1.00,12),
(20,'Machining','EDB05','Máy dán cạnh thẳng tự động','Lỗi nồi keo',T(18,0),T(21,0),3.00,12),
(20,'Machining','CNCR09','Máy chép hình thẳng','Lỗi trục không quay',T(17,50),T(19,50),2.00,12),
(21,'RC','VEC04','Máy cắt veneer','Máy tự cắt',T(18,35),T(20,0),1.42,12),
(22,'Machining','CNCR20','Máy CNC','Gãy giá đỡ dây điện',T(16,49),T(17,40),0.85,12),
(24,'Machining','CNCR11','Hogmag','Lỗi hút chân không',T(7,40),T(7,53),0.22,12),
(25,'RC','VC14','Máy hút chân không','Lỗi hút chân không',T(8,25),T(10,9),1.73,12),
]
rows=[]
for d,ar,mc,mn,fail,fr,to,h,n in rep:
    notes='-'.join(x for x in (mc,mn,fail) if x)
    rows.append((D(d),'A10',notes,{'RC':'Raw Cutting'}.get(ar,ar),'MAINTENANCE',fr,to,None))
wb=openpyxl.Workbook();wb.remove(wb.active)
B.ORDER.insert(0,'A10')
for a in ('Machining','Raw Cutting'):
    ws=B.sheet(wb,a,[x for x in rows if x[3]==a])
    ws['A2']='WORKER QTY ĐỂ TRỐNG: chỉ điền SỐ CÔNG NHÂN THẬT đứng máy khi có xác nhận của tổ trưởng khu vực (ghi tên người xác nhận ở cột K). Không dùng cột số cuối (12/24) của báo cáo bảo trì. '+ws['A2'].value
    ws.cell(4,11,'Người xác nhận số công nhân').fill=B.gry;ws.cell(4,11).font=Font(name=B.F,bold=True);ws.cell(4,11).alignment=B.ctr;ws.cell(4,11).border=B.bd;ws.column_dimensions['K'].width=26
    for r_ in range(5,ws.max_row+1):
        if isinstance(ws.cell(r_,1).value,dt.datetime): ws.cell(r_,8).fill=PatternFill('solid',fgColor='FFFF00');ws.cell(r_,11).border=B.bd
    ws['A2'].font=Font(name=B.F,italic=True,color='C00000')
F=B.F;bd=B.bd
ws=wb.create_sheet('Đối chiếu báo cáo bảo trì')
H=['Ngày','Khu vực (báo cáo)','Mã máy','Tên máy','Hiện tượng','Từ','Đến','Giờ dừng báo cáo','Giờ dừng tính lại','Chênh lệch','Cột số cuối của báo cáo (chưa rõ nghĩa, KHÔNG dùng làm số công nhân)','Giờ công (chưa tính, cần số công nhân thật)']
for i,h in enumerate(H,1):
    c=ws.cell(2,i,h);c.fill=B.red;c.font=B.hf;c.alignment=B.ctr;c.border=bd
ws['A1']='ĐỐI CHIẾU VỚI BÁO CÁO BẢO TRÌ (Mr.Phong - Mr.Khiêm, Maint) 19/10-25/10';ws['A1'].font=Font(name=F,bold=True,size=13)
for k,(d,ar,mc,mn,fail,fr,to,h,n) in enumerate(rep):
    r=3+k
    for c,v in enumerate([D(d),ar,mc,mn,fail,fr,to,h,f'=(G{r}-F{r})*24',f'=ROUND(I{r}-H{r},2)',n,None],1):
        x=ws.cell(r,c,v);x.font=Font(name=F);x.border=bd;x.alignment=Alignment(horizontal='left' if c in(4,5) else 'center',vertical='center',wrap_text=c in(4,5))
    ws.cell(r,1).number_format='dd/mm/yyyy';ws.cell(r,6).number_format=ws.cell(r,7).number_format='hh:mm';ws.cell(r,9).number_format='0.00';ws.cell(r,12).number_format='0.0'
e=2+len(rep)
ws.cell(e+1,8,'Tổng giờ dừng máy').font=Font(name=F,bold=True);ws.cell(e+1,9,f'=SUM(I3:I{e})').font=Font(name=F,bold=True);ws.cell(e+1,9).number_format='0.00'
ws.cell(e+3,1,'Ghi chú: 20/10 máy bơm sơn chuyền UV không có mã máy trong báo cáo. 25/10 là Chủ nhật, báo cáo ghi sự cố máy VC14 khu vực RC (cần xác nhận Raw Cutting có làm hôm đó không). EDB05 ngày 20/10 dừng đến 21:00 (sau giờ tăng ca 20:00).').font=Font(name=F,italic=True)
for col,w in zip('ABCDEFGHIJKL',[12,12,10,28,30,8,8,12,12,10,14,18]): ws.column_dimensions[col].width=w
wb.save('Nhap_lieu_A10_19-10_den_25-10_NHAP.xlsx')
