import openpyxl,datetime as dt,collections
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter as L
Z='s32zip/Lost time S32/'
D0=dt.datetime(2026,10,5)
T=lambda h,m:dt.time(h,m)
def ld(f,sheet): return list(openpyxl.load_workbook(Z+f,read_only=True,data_only=True)[sheet].iter_rows(values_only=True))
R=[]   # dict rows
DM={'Finishing':'FINISHING','P1':'OTHERS (INVESTIGATE)','TF':'TECHNICAL FINISHING','ME':'ME','QC':'QA/QC/TESTING','Machining':'MACHINING','Subcon':'SUBCON'}
AM={'RC - PANEL':'Raw Cutting'}
def add(date,src,item,qty,area,dept,lido,gc,fr,to,hc,flag='',items=None,chk=None):
    dept=DM.get(dept,dept);area=AM.get(area,area)
    R.append(dict(date=date,src=src,item=item,qty=qty,area=area,dept=dept,lido=lido,gc=gc,fr=fr,to=to,hc=hc,flag=flag,items=items or []))
# 1. Beam saw (bản 20/10 mới nhất)
bs=[r for r in ld('LOST TIME BEMSAW 24.10.xlsx','data ') if isinstance(r[6],dt.datetime) and r[6]>=D0]
assert len(bs)==24
notime=[r for r in bs if r[8] is None];bs=[r for r in bs if r[8] is not None]
print('Beam saw dòng không có giờ:',[(r[6].strftime('%d/%m'),r[2],r[3],r[5]) for r in notime])
dm={'HÀNG BẢNG MÀU':'TF','FINISHING':'Finishing','Machining CNC':'Machining','Machining':'Machining','Unassigned':'(chưa phân công)','QA do không test lú':'QC'}
for r in bs:
    ch=str(r[13]).strip()
    dept='QC' if ch.upper().startswith('QA') else dm[ch]
    flag='U' if ch=='Unassigned' else ''
    gc=f'Charge to: {ch}' if (ch in('Unassigned','HÀNG BẢNG MÀU') or ch.upper().startswith('QA')) else ''
    fr_,to_=r[8],r[9]
    if r[10]<0:
        flag='E';gc=(gc+'; ' if gc else '')+f'Nguồn ghi {fr_:%H:%M}-{to_:%H:%M} (giờ kết thúc nhỏ hơn giờ bắt đầu, -0.1 giờ)';to_=fr_
    add(r[6],'Beam saw',r[2],r[5],'RC - PANEL',dept,f'Beam saw - phát sinh (chi tiết {str(r[3]).strip()})',gc or f'STT {r[0]}',fr_,to_,int(r[7]),flag)
# 2. UV hằng ngày (W41 từ 5/10): hàng làm việc 8h + OT 4h
uvraw=[r for r in ld('lost time UV 24.10.xlsx','LOST TIME') if isinstance(r[1],dt.datetime) and r[1]>=D0]
uv=[];seen=set()
for r in uvraw:
    if r[1] in seen: assert any(tuple(x[:6])==tuple(r[:6]) for x in uv);continue
    seen.add(r[1]);uv.append(r)
print('UV báo cáo ngày:',len(uvraw),'dòng (trùng',len(uvraw)-len(uv),') ->',len(uv),'ngày')
for r in uv:
    d=r[1];work=int(r[3] or 0);ot=int(r[4] or 0)
    if work: add(d,'UV hằng ngày',None,None,'RC - PANEL','Finishing','Không có bomtime chuyền UV','HC làm việc (ca thường)',T(7,30),T(15,30),work,'U' if d.weekday()==6 else '')
    if ot: add(d,'UV hằng ngày',None,None,'RC - PANEL','Finishing','Không có bomtime chuyền UV','HC tăng ca (OT 4 giờ)',T(15,30),T(19,30),ot,'')
# 3. UV tổng hợp phát sinh (Sheet1 bản (1))
s1=[r for r in ld('lost time UV 24.10.xlsx','Sheet1')[2:] if r[0] is not None]
g12=[r for r in s1 if r[6]==dt.datetime(2026,10,12)]
items12=[(r[2],r[3],r[5]) for r in g12]
hrs=[r for r in g12 if r[7] is not None]
assert len(g12)==23 and len(hrs)==2
for k,r in enumerate(hrs):
    add(dt.datetime(2026,10,12),'UV phát sinh',g12[0][2] if k==0 else None,None,'RC - PANEL','QC','Phát sinh khu vực UV - SLAT (23 mã hàng)',f'Charge to: QC A HIẾU; mã hàng xem sheet Chi tiết',r[8],r[9],int(r[7]),'',items12 if k==0 else [])
g13=[r for r in s1 if r[6]==dt.datetime(2026,10,13)]
assert len(g13)==1
r=g13[0];add(r[6],'UV phát sinh',r[2],r[5],'RC - PANEL','TF','Phát sinh khu vực UV - Back Panel','Charge to: TF',r[8],r[9],int(r[7]))
# 4. Ép nóng (bản (1))
ep=ld('FILE ÉP HÀNG SỬA ( KV ÉP NÓNG ) (1).xlsx','datta')
def grp(i0,i1): return [(ep[i-1][2],ep[i-1][3],ep[i-1][5]) for i in range(i0,i1+1)]
def eprow(i): return ep[i-1]
r=eprow(219);add(r[6],'Ép nóng',r[2],None,'RC - PANEL','P1','Hỗ trợ P1 ép nóng','Hỗ trợ P1 (3 mã hàng, 1 ca)',r[8],r[9],int(r[7]),'',grp(219,221))
r=eprow(222);add(r[6],'Ép nóng',r[2],None,'RC - PANEL','P1','Hỗ trợ P1 ép nóng - ép lót','Ép lót (2 mã hàng)',r[8],r[9],int(r[7]),'',grp(222,223))
r=eprow(224);add(r[6],'Ép nóng',r[2],None,'RC - PANEL','P1','Hỗ trợ P1 ép nóng - ép chính','Ép chính (2 mã hàng)',r[8],r[9],int(r[7]),'',grp(224,225))
# 5. Veneer hàng BS (bản 15/10)
ve=ld('Lost time veneer 24.10b.xlsx','tổng')
v1,v2,v3,v4,v5,v6,v7,v8,v9=[ve[i] for i in range(1802,1811)]
add(v1[11],'Veneer',None,None,'RC - PANEL','TF','Làm mẫu (veneer)','2 mã (code 516, code 231), 1 ca',v1[8],v1[9],int(v1[7]),'',[(v1[1],v1[3],v1[5]),(v2[1],v2[3],v2[5])])
add(v3[11],'Veneer',None,None,'RC - PANEL','ME','Làm mẫu (veneer)','code 229',v3[8],v3[9],int(v3[7]),'',[(v3[1],v3[3],v3[5])])
add(v4[11],'Veneer',None,None,'RC - PANEL','TF','Làm mẫu (veneer)','code 191',v4[8],v4[9],int(v4[7]),'',[(v4[1],v4[3],v4[5])])
add(v5[11],'Veneer',None,None,'RC - PANEL','(cần xác nhận)','Veneer walnut cắt ghép (bs lại)','Nguồn ghi Charge to: "bs lại"',v5[8],v5[9],int(v5[7]),'E',[(v5[2],v5[3],v5[5])])
add(v6[11],'Veneer',None,None,'RC - PANEL','ME','Làm mẫu (veneer)','code 014',v6[8],v6[9],int(v6[7]),'',[(v6[1],v6[3],v6[5])])
add(v7[11],'Veneer',None,None,'RC - PANEL','ME','Soạn hàng veneer lỗi trả lại kho','code 239 (130 m2)',v7[8],v7[9],int(v7[7]),'',[(v7[1],v7[3],v7[5])])
add(v8[11],'Veneer','RH12-4706-0423',v8[5],'RC - PANEL','ME','Sửa hàng hộc kéo (veneer)','Hàng hộc kéo, 4 người cả ngày',v8[8],v8[9],int(v8[7]),'')
add(v9[11],'Veneer','RH12-4706-0423',v9[5],'RC - PANEL','ME','Sửa hàng hộc kéo (veneer)','Hàng hộc kéo, 2 người',v9[8],v9[9],int(v9[7]),'')
# 5b. Hàng sửa khu vực MAC (Mr.Hiệp)
mc=openpyxl.load_workbook(Z+'hang sua mac.b.xlsx',data_only=True)
m2=[r for r in mc['Sheet2'].iter_rows(values_only=True) if isinstance(r[6],dt.datetime) and r[6]>=D0][0]
m3=[r for r in mc['Sheet3'].iter_rows(values_only=True) if isinstance(r[6],dt.datetime) and r[6]>=D0][0]
add(m2[6],'MAC',m2[1],None,'Machining','ME','Hỗ trợ sửa hàng MAC - chạy nhỏ lại mộng theo routing',f'Chi tiết {m2[2]}, 100 cây',m2[7],m2[8],int(m2[11]),'',[(m2[1],m2[2],100)])
add(m3[6],'MAC',m3[1],None,'Machining','ME','Hỗ trợ sửa hàng MAC - khoan lại lỗ vít 11mm',f'Chi tiết {m3[2]}, 65 bộ. Nguồn ghi tổng 1:55 nhưng 9:32-11:20 là 1:48',m3[7],m3[8],int(m3[11]),'P',[(m3[1],m3[2],65)])
# 6. Ba hình (Assembly)
for d in range(12,17): add(dt.datetime(2026,10,d),'Thợ sơn (hình)',None,None,'Assembly','Finishing','Không có bomtime sơn lòng trong','Thợ sơn lòng trong tủ - hỗ trợ Fin',T(7,30),T(19,30),12)
add(dt.datetime(2026,10,17),'Thợ sơn (hình)',None,None,'Assembly','Finishing','Không có bomtime sơn lòng trong','Thợ sơn lòng trong tủ - hỗ trợ Fin (thứ Bảy)',T(7,30),T(15,30),12,'P')
add(dt.datetime(2026,10,16),'Cào cước (hình)','RH893',None,'Assembly','Machining','Cào cước hỗ trợ MAC','Line 8 - hỗ trợ MAC',T(15,0),T(20,0),1,'P')
add(dt.datetime(2026,10,16),'Cào cước (hình)','RH5061',None,'Assembly','Machining','Cào cước hỗ trợ MAC','Line 7 - hỗ trợ MAC',T(8,0),T(12,0),1,'')
add(dt.datetime(2026,10,15),'Bắn cát (hình)','RH2252',734,'Assembly','Subcon','Bắn cát hỗ trợ Subcon Đại Xuyên (chân)','Line bắn cát - hỗ trợ Supcon Đại Xuyên',T(7,0),T(15,0),4,'A')
order=['Ép nóng','UV hằng ngày','Beam saw','Veneer','MAC','UV phát sinh','Thợ sơn (hình)','Cào cước (hình)','Bắn cát (hình)']
R.sort(key=lambda x:(x['date'],order.index(x['src']),x['fr']))
print('số dòng',len(R))
import pickle;pickle.dump(R,open('s32all.pkl','wb'))
hours=lambda x:((x['to'].hour*60+x['to'].minute)-(x['fr'].hour*60+x['fr'].minute))/60*x['hc']
by=collections.defaultdict(float)
for x in R: by[x['src']]+=hours(x)
for k,v in by.items(): print(k,round(v,3))
print('tổng',round(sum(by.values()),3))
